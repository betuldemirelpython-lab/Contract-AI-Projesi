"""
contract_service.py - Sözleşme İşleme Servisi
PDF, DOCX okuma ve sözleşme işleme pipeline'ı
"""

import io
import logging
import time
from pathlib import Path
from typing import Optional, BinaryIO, Union

logger = logging.getLogger(__name__)

UPLOADS_DIR = Path("uploads")
UPLOADS_DIR.mkdir(exist_ok=True)


# ─── PDF Metin Çıkarma ────────────────────────────────────────────────────
def extract_text_from_pdf(file_content: bytes, filename: str = "dosya.pdf") -> str:
    """
    PyMuPDF ile PDF'ten metin çıkarır.
    file_content: raw bytes
    """
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(stream=file_content, filetype="pdf")
        text_parts = []

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text("text")
            if text.strip():
                text_parts.append(f"[Sayfa {page_num + 1}]\n{text}")

        doc.close()
        full_text = "\n\n".join(text_parts)

        if not full_text.strip():
            raise ValueError("PDF'ten metin çıkarılamadı. Taranmış PDF olabilir.")

        logger.info(f"PDF okundu: {filename}, {len(full_text)} karakter")
        return full_text

    except ImportError:
        raise ImportError("PyMuPDF kurulu değil. `pip install pymupdf` çalıştırın.")
    except Exception as e:
        raise RuntimeError(f"PDF okuma hatası ({filename}): {e}")


# ─── DOCX Metin Çıkarma ───────────────────────────────────────────────────
def extract_text_from_docx(file_content: bytes, filename: str = "dosya.docx") -> str:
    """
    python-docx ile DOCX'ten metin çıkarır.
    """
    try:
        from docx import Document

        doc = Document(io.BytesIO(file_content))
        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Tablolar da dahil et
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n".join(text_parts)

        if not full_text.strip():
            raise ValueError("DOCX'ten metin çıkarılamadı.")

        logger.info(f"DOCX okundu: {filename}, {len(full_text)} karakter")
        return full_text

    except ImportError:
        raise ImportError("python-docx kurulu değil. `pip install python-docx` çalıştırın.")
    except Exception as e:
        error_msg = str(e)
        if "is not a Word file" in error_msg or "not a zip file" in error_msg.lower():
            if filename.lower().endswith(".doc"):
                raise RuntimeError(
                    f"Eski format (.doc) veya geçersiz dosya ({filename}). "
                    "Sistem sadece güncel .docx formatını desteklemektedir. "
                    "Lütfen dosyanızı Word'de açıp '.docx' formatında 'Farklı Kaydet' diyerek tekrar yükleyin."
                )
            else:
                raise RuntimeError(
                    f"Dosya bozuk veya geçerli bir Word dosyası değil ({filename}). Lütfen dosyayı kontrol edip tekrar yükleyin."
                )
        raise RuntimeError(f"DOCX okuma hatası ({filename}): {e}")


# ─── TXT Metin Çıkarma ────────────────────────────────────────────────────
def extract_text_from_txt(file_content: bytes, filename: str = "dosya.txt") -> str:
    """Plain text dosyasını okur."""
    for encoding in ["utf-8", "cp1254", "latin-1", "iso-8859-9"]:
        try:
            text = file_content.decode(encoding)
            logger.info(f"TXT okundu ({encoding}): {filename}")
            return text
        except (UnicodeDecodeError, LookupError):
            continue
    raise RuntimeError(f"Dosya okunamadı ({filename}): Desteklenmeyen karakter kodlaması")


# ─── Dosya Türüne Göre Metin Çıkar ───────────────────────────────────────
def extract_text_from_file(file_content: bytes, filename: str) -> str:
    """
    Dosya uzantısına göre doğru parser'ı kullanır.
    """
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(file_content, filename)
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(file_content, filename)
    elif suffix in (".txt", ".text"):
        return extract_text_from_txt(file_content, filename)
    else:
        raise ValueError(
            f"Desteklenmeyen dosya türü: {suffix}. "
            "PDF, DOCX veya TXT yükleyin."
        )


# ─── Dosyayı Kaydet ───────────────────────────────────────────────────────
def save_uploaded_file(file_content: bytes, filename: str) -> Path:
    """Yüklenen dosyayı uploads/ dizinine kaydeder."""
    timestamp = int(time.time())
    safe_name = f"{timestamp}_{filename}"
    dest = UPLOADS_DIR / safe_name
    dest.write_bytes(file_content)
    logger.info(f"Dosya kaydedildi: {dest}")
    return dest


# ─── Metin Doğrulama ──────────────────────────────────────────────────────
def validate_contract_text(text: str) -> tuple[bool, str]:
    """
    Sözleşme metninin analiz için uygun olup olmadığını kontrol eder.
    Returns: (geçerli_mi, hata_mesajı)
    """
    if not text or not text.strip():
        return False, "Sözleşme metni boş olamaz."

    word_count = len(text.split())
    if word_count < 20:
        return False, f"Sözleşme metni çok kısa ({word_count} kelime). En az 20 kelime olmalı."

    if len(text) > 200_000:
        return False, "Sözleşme metni çok uzun (200.000 karakter sınırı). Lütfen kısaltın."

    return True, ""


# ─── Metin Önizleme ───────────────────────────────────────────────────────
def get_text_preview(text: str, max_chars: int = 300) -> str:
    """Metnin ilk N karakterini döndürür."""
    preview = text[:max_chars].strip()
    if len(text) > max_chars:
        preview += "..."
    return preview


# ─── İstatistikler ────────────────────────────────────────────────────────
def get_text_stats(text: str) -> dict:
    """Metin hakkında temel istatistikler döndürür."""
    words = text.split()
    sentences = text.count(".") + text.count("!") + text.count("?")
    return {
        "karakter_sayisi": len(text),
        "kelime_sayisi": len(words),
        "satir_sayisi": text.count("\n"),
        "cumle_sayisi": max(sentences, 1),
        "tahmini_okuma_suresi": f"{max(1, len(words) // 200)} dakika",
    }
